import csv
import os
import pyodbc
import smtplib
import pandas as pd,re
import random,string
import json
from datetime import datetime, timedelta
import time
import shutil
from multiprocessing import Queue
import pickle
from email.mime.text import MIMEText
import openpyxl
from openpyxl import Workbook
from email.mime.multipart import MIMEMultipart 
from email.mime.application import MIMEApplication
from .send_data_to_sql import generate_edi_segment_excel
from .tests import insert_data_to_DB2
import sqlite3
from datetime import datetime
from email.mime.base import MIMEBase
from email import encoders
from io import BytesIO

# Email Configuration
smtp_config = {
    'host': 'mail.privateemail.com',
    'port': 465,
    'user': 'support@disruptionsim.com',
    'password': 'Onesmarter@2023'
}

csv_headers = [
    "SUB/DEP", "LAST NAME", "FIRST NAME", "SSN","TEMP SSN","SEX", "DOB", "DEP LAST NAME", "DEP FIRST NAME",
    "DEP SSN", "DEP SEX", "DEP DOB","CUSTODIAL PARENT","LOCAL", "PLAN", "CLASS", "EFF DATE", "TERM DATE", "ID",
    "ADDRESS 1", "ADDRESS 2", "CITY", "STATE", "ZIP", "PHONE", "EMAIL", "STATUS", "TYPE","MEMBER ID","DEP ADDRESS","DEP CITY","DEP STATE","DEP ZIP"
]


def send_success_email_count(emails, count_list, term_ssn_list):
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['Subject'] = "Count and term data for xlsx and db2"

    count_table = """
    <table border="1" cellpadding="5" cellspacing="0">
        <tr><th>Source</th><th>Total Subscribers</th><th>Total Spouse</th><th>Other Dependents</th></tr>
        <tr><td>Excel</td><td>{}</td><td>{}</td><td>{}</td></tr>
        <tr><td>DB2</td><td>{}</td><td>{}</td><td>{}</td></tr>
    </table>
    """.format(
        count_list[0]['excel'][0]['total_subscriber'], 
        count_list[0]['excel'][1]['total_spouse'], 
        count_list[0]['excel'][2]['other dependents'],
        count_list[1]['db2'][0]['total_subscriber'], 
        count_list[1]['db2'][1]['total_spouse'], 
        count_list[1]['db2'][2]['other dependents']
    )

    term_table = ""
    if term_ssn_list:
        term_table = "<br/><br/><strong>Terminated Employees:</strong><table border='1' cellpadding='5' cellspacing='0'><tr><th>SSN</th><th>TERM DATE</th></tr>"
        for entry in term_ssn_list:
            term_table += "<tr><td>{}</td><td>{}</td></tr>".format(entry['SSN'], entry['TERM DATE'])
        term_table += "</table>"

    body = f"""
    <p>Count list for both Excel and DB2:</p>
    {count_table}
    {term_table}
    """

    msg.attach(MIMEText(body, 'html'))

    for email in emails:
        msg['To'] = email
        server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
        print(f"Success email sent to {email}.")

    server.quit()

def send_success_email_term(emails,count_list):
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['Subject'] = f"SSN with term date"

    body = f"""
    <p>These are the ssn added to db2 and excel with a Term date {count_list}.</p>
    """
    msg.attach(MIMEText(body, 'html'))
    
    for email in emails:
        msg['To'] = email
        server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
        print(f"Success email sent to {email}.")

def send_success_email_count_everything_fine(emails,count_list,term_ssn_list):
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['Subject'] = f"Count matched for xlsx and db2"

    body = f"""
    <p>Count list for both excel and db2 {count_list}.</p>
    """
    msg.attach(MIMEText(body, 'html'))
    
    for email in emails:
        msg['To'] = email
        server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
        print(f"Success email sent to {email}.")



def send_success_email(email, file_name, output_path):
    
                        
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['To'] = email
    msg['Subject'] = f"Processing Successful: {file_name}"

    body = f"""
    <p>The file <strong>{file_name}</strong> was processed successfully.</p>
    <p>Please find the processed file attached.</p>
    """
    msg.attach(MIMEText(body, 'html'))
    
    if file_name.endswith('.X12'):
                print('x12')
                file_name = file_name.replace('.X12', '.csv')
                print('csv')

    # Attach the processed output file
    with open(output_path, 'rb') as f:
        part = MIMEApplication(f.read(), Name="file_name")
        part['Content-Disposition'] = f'attachment; filename="{file_name}"'
        msg.attach(part)

    server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
    server.quit()
    print(f"Success email sent for {file_name} to {email}")

def send_error_email(email, file_name, error_message):
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['To'] = email
    msg['Subject'] = f"Processing Failed: {file_name}"

    body = f"""
    <p>The file <strong>{file_name}</strong> failed to process.</p>
    <p><strong>Reason:</strong> {error_message}</p>
    """
    msg.attach(MIMEText(body, 'html'))

    server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
    server.quit()
    print(f"Error email sent for {file_name} to {email}")


def send_error_email_variance(email,file_name,p_type):
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['To'] = email
    msg['Subject'] = f"Processing Failed: {file_name}"

    body = f"""
    <p>The file <strong>{file_name}</strong> has a variance in {p_type} of more than five percentage  .</p>
    """
    msg.attach(MIMEText(body, 'html'))

    server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
    server.quit()
    print(f"Error email sent for {file_name} to {email}")

def send_member_id_email(new_ids, missing_ids,email):
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['To'] = email
    msg['Subject'] = f"Processing Successful"

    # HTML body for a well-structured email
    body = f"""
    <html>
    <body>
        <p>Hello,</p>

        <p>Here is the report for Member IDs:</p>

        <h3>New Member IDs (present today, not in previous data):</h3>
        <p>{', '.join(new_ids) if new_ids else 'None'}</p>

        <h3>Missing Member IDs (present in previous data, not today):</h3>
        <p>{', '.join(missing_ids) if missing_ids else 'None'}</p>

        <br>
        <p>Best regards,</p>
        <p>Your System</p>
    </body>
    </html>
    """
    msg.attach(MIMEText(body, "html"))
    server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
    server.quit()
    print(f"Success email sent to {email}")


def send_error_log_email(email, file_name, error_message, error_logs):
    # Step 1: Convert error_logs to DataFrame
    data = [{"Member ID": key, "Group Number": value} for key, value in error_logs.items()]
    df = pd.DataFrame(data)

    # Step 2: Save DataFrame to an in-memory Excel file
    excel_buffer = BytesIO()
    with pd.ExcelWriter(excel_buffer, engine='xlsxwriter') as writer:
        df.to_excel(writer, index=False, sheet_name='Error Logs')
    excel_buffer.seek(0)

    # Step 3: Prepare the email content
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['To'] = ", ".join(email)
    msg['Subject'] = f"Group Numbers Not Found: {file_name}"

    # Email body
    body = f"""
    <p><strong>Group numbers not found for the following member ID(s):</strong> {error_message}</p>
    """
    msg.attach(MIMEText(body, 'html'))

    # Step 4: Attach the Excel file
    part = MIMEBase('application', 'octet-stream')
    part.set_payload(excel_buffer.read())
    encoders.encode_base64(part)
    part.add_header(
        'Content-Disposition',
        f'attachment; filename="ErrorLogs_{file_name}.xlsx"'
    )
    msg.attach(part)

    # Step 5: Send the email
    server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
    server.quit()
    print(f"Group number log email sent for {file_name} to {email}")

def send_missing_ssn_email(missing_ssn,missing_in_inventory,email):
    server = smtplib.SMTP_SSL(smtp_config['host'], smtp_config['port'])
    server.login(smtp_config['user'], smtp_config['password'])

    msg = MIMEMultipart()
    msg['From'] = smtp_config['user']
    msg['To'] = email
    msg['Subject'] = f"Processing Successful"

    # HTML body for a well-structured email
    body = f"""
    <html>
    <body>
        <p>Hello,</p>

        <p>Here is the report for Member IDs:</p>

        <h3>Missing ssn (present in inventory, not in current data):</h3>
        <p>{', '.join(missing_ssn) if missing_ssn else 'None'}</p>

        <h3>Missing inventory ssn (present in current data, not in inventory):</h3>
        <p>{', '.join(missing_in_inventory) if missing_in_inventory else 'None'}</p>

        <br>
        <p>Best regards,</p>
        <p>Your System</p>
    </body>
    </html>
    """
    msg.attach(MIMEText(body, "html"))
    server.send_message(msg, from_addr=smtp_config['user'], to_addrs=email)
    server.quit()
    print(f"Success email sent to {email}")


def parse_edi_to_csv(input_file_path, output_directory,system_directory):
    os.makedirs(output_directory, exist_ok=True)
    os.makedirs(system_directory, exist_ok=True)
    output_csv_path = os.path.join(output_directory, os.path.basename(input_file_path))
    system_csv_path = os.path.join(system_directory, os.path.basename(input_file_path))
    file_name = input_file_path.split("/")[-1]  # Extracts 'EDI_834_11-15-2024_3KXST5r.X12'
    #date_part = file_name.split("_")[2]
    #print("Extracted Date:", date_part)
    #date_part = str(date_part[:10])
    date_part='NA'
    with open(input_file_path, 'r') as file:
        edi_data = file.read()
    segments = edi_data.strip().split("~")
    csv_data = []
    current_subscriber = {}
    dependents = []
    error_logs = {}
    segment_list = []
    parsed_data_list = []
    total_parsed_data = []
    cus_data_list = []
    k = 0
    def extract_segment_data(segment, delimiter="*"):
        return segment.split(delimiter)

    for segment in segments:
        each_segments = segment.split("*") 
        segment_name = each_segments[0]  
        parsed_data = {}
        if segment_name in ["ISA", "GS", "ST", "BGN", "DTP", "N1", "INS", "REF", "NM1", 
                            "PER", "N3", "N4", "DMG", "HD", "SE", "GE", "IEA"]:
            parsed_data[segment_name] = "*".join(each_segments[1:])
            parsed_data_list.append(parsed_data)
            if segment_name == "HD":
                total_parsed_data.append({k:parsed_data_list})
                k += 1
                parsed_data_list = []
        else:
            print(f"Skipping unknown segment: {segment_name}")
        
        elements = extract_segment_data(segment)
        segment_id = elements[0]
        if segment_id not in segment_list:
            segment_list.append(segment_id)
        if segment_id == "REF":
            member_id_code = elements[1]
            if(member_id_code == "0F"):
                member_id = elements[2]
        if segment_id == "INS":
            relationship_code = elements[2]
            if relationship_code == '18':
                Sub = "Subscriber"
                Type = '18'
            else:
                dependent_type = elements[2]
                if dependent_type == '01':
                    Sub = "Spouse"
                    Type= dependent_type
                elif dependent_type == '19':
                    Sub = "Child"
                    Type = dependent_type
                else:
                    Sub = "Dependent"
                    Type= dependent_type
            if elements[1] == 'Y':
                status = 'Active'
            elif elements[1] == 'N':
                status = 'Inactive'
            else:
                status = ''

        elif segment_id == "NM1" and elements[1] == "IL":
            if current_subscriber:
                csv_data.append(current_subscriber)
                current_subscriber = {}
            sss = elements[-1] if len(elements) > 8 else ""
            sss = sss.replace("-", "").strip()
            if len(sss) == 9:
                sss = f"{sss[:3]}-{sss[3:5]}-{sss[5:]}"
            elif len(sss) == 8:
                sss = f"{sss[:2]}-{sss[2:4]}-{sss[4:]}"
            else:
                sss = "" 
            person = {
                "LAST NAME": elements[3] if len(elements) > 3 else "",
                "FIRST NAME": elements[4] if len(elements) > 4 else "",
                "SSN": sss,
                "SUB/DEP": Sub,
                "STATUS":status,
                "TYPE":Type,
                "MEMBER ID": member_id
            }
            current_subscriber.update(person)

        elif segment_id == "DMG" and len(elements) > 2:
            dob = elements[2]
            person = dependents[-1] if dependents else current_subscriber
            person["DOB"] = f"{dob[4:6]}/{dob[6:]}/{dob[:4]}" if len(dob) == 8 else ""
            person["SEX"] = elements[3] if len(elements) > 3 else ""
        
        elif "REF*17" in segment:
            data = segment.split("*")
            cus_data = data[-1]
            person["CUSTODIAL PARENT"] = cus_data

        elif segment_id == "N3" and len(elements) > 1:
            address = elements[1]
            person = dependents[-1] if dependents else current_subscriber
            person["ADDRESS 1"] = address

        elif segment_id == "N4" and len(elements) > 3:
            city, state, zip_code = elements[1:4]
            zerocode = zip_code.zfill(5)
            zip_code = str(zip_code).strip()
            if len(zip_code) < 5:
                zip_code = zip_code.zfill(5)
            elif len(zip_code) > 5:
                zip_code = zip_code[:5] 
            person = dependents[-1] if dependents else current_subscriber
            person.update({"CITY": city, "STATE": state, "ZIP": str(zip_code)})
        elif segment_id == "PER" and len(elements) > 3:
            phone = elements[-1]
            person = dependents[-1] if dependents else current_subscriber
            person["PHONE"] = phone

        elif segment_id == "HD" and len(elements) > 2:
            current_subscriber["PLAN"] = elements[1]
            current_subscriber["CLASS"] = elements[3] if len(elements) > 3 else ""

        elif segment_id == "DTP" and len(elements) > 2:
            if elements[1] == "348":
                eff_date = elements[-1]
                current_subscriber["EFF DATE"] = f"{eff_date[4:6]}/{eff_date[6:]}/{eff_date[:4]}" if len(eff_date) == 8 else ""
            elif elements[1] == "349":
                term_date = elements[-1]
                current_subscriber["TERM DATE"] = f"{term_date[:4]}/{term_date[4:6]}/{term_date[6:]}" if len(term_date) == 8 else ""

        elif segment_id == "REF" and len(elements) > 2 and elements[1] == "1L":
            current_subscriber["ID"] = elements[2]
            if elements[2] == "L11958M001":
                current_subscriber["PLAN"] = str("01")
                current_subscriber["CLASS"] = "01"
            
            elif elements[2] == "L11958M002":
                current_subscriber["PLAN"] = str("01")
                current_subscriber["CLASS"] = "02"
                
            elif elements[2] == "L11958MD01":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "SS"
                
            elif elements[2] == "L11958MR01":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "R8"
                
            elif elements[2] == "L11958MR02":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "D9"
                
            elif elements[2] == "L11958MR03":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "R1"    
                
            elif elements[2] == "L11958MR04":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "D2"       
                
            elif elements[2] == "L11958MR05":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "M8"             
                
            elif elements[2] == "L11958MR06":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "M9"    
                
            elif elements[2] == "L11958MR07":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "M1"
                
            elif elements[2] == "L11958MR08":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "M2"   
                
            elif elements[2] == "L11958MR09":
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "D0"     
                
            else:
                current_subscriber["PLAN"] = "01"
                current_subscriber["CLASS"] = "01"

                error_logs[member_id] = elements[2]   
    errorFileName =  os.path.basename(input_file_path)
    # print(errorFileName)
    # if(len(error_logs) >0):
    #     error_message = "Missing group numbers for the given Member IDs"
    #     email = ['krishnarajjadhav2003@gmail.com']
    #     send_error_log_email(email, errorFileName, error_message, error_logs) 

    conn = sqlite3.connect("db.sqlite3")
    cursor = conn.cursor()
    import pandas as pd
    flattened_data = []
    flattened_data = []
    for group in total_parsed_data:
        for group_id, records in group.items():
            for record in records:
                for key, value in record.items():
                    flattened_data.append({'group_id': group_id, 'key': key, 'value': value})

    df = pd.DataFrame(flattened_data)
    df = df.groupby(['group_id', 'key'], as_index=False).agg({'value': 'first'})

    pivot_df = df.pivot(index='group_id', columns='key', values='value').reset_index()
    pivot_df = pivot_df.where(pd.notnull(pivot_df), None)
    pivot_df.drop(columns=['group_id'],inplace=True)
    for column in pivot_df.columns:
        pivot_df[column] = pivot_df[column].str.replace('*', '  ', regex=False)
        pivot_df[column] = pivot_df[column].drop_duplicates().reset_index(drop=True)
    pivot_df = pivot_df.fillna(' ')
    pivot_df['Date_edi'] = date_part
    random_number = random.randint(0, 9999)
    random_alphabet = random.choice(string.ascii_uppercase) 
    result = f"{random_alphabet}{random_number:04}"
    out_dir = "media/csv_files/"
    pivot_df_data = pivot_df.to_dict(orient="records")
    segment_df = generate_edi_segment_excel(pivot_df_data)
    try:
        edi_excel_path = os.path.join(out_dir, f"edi_segment_data_{result}.xlsx")
        segment_df.to_excel(edi_excel_path)
    except:
        edi_excel_path = "S:\\OOE\\EDI_PROJECT-\\EDI-Backend\\media\\csv_files"
        new_path = os.path.join(edi_excel_path,f"edi_segment_data_{result}.xlsx")
        print(new_path)
        segment_df.to_excel(new_path)
        edi_excel_path = new_path
    def write_to_queue(data, queue_file):
        try:
            with open(queue_file, "ab") as file:
                pickle.dump(data, file)
            print("Data added to the queue.")
        except Exception as e:
            print(f"Error writing to the queue: {e}")
    write_to_queue(pivot_df_data, "queue_file.pkl")

    # send_data_to_serever(pivot_df_data)
    # send_data_in_json_form(pivot_df_data)
    conn.close()
    csv_data.append(current_subscriber)
    csv_data.extend(dependents)
    input_filename = os.path.splitext(os.path.basename(input_file_path))[0]
    output_csv_path = os.path.join(output_directory, f"{input_filename}.csv")
    output_xlsx_path = os.path.join(out_dir, f"{input_filename}.xlsx")
    system_csv_path = os.path.join(system_directory, f"{input_filename}.csv")
    for row in csv_data:
        if 'ID' in row.keys():
            row['STATUS'] = row['ID']
        else:
            row['STATUS'] = ''
        if 'TYPE' in row.keys():
            row['ID'] = row['TYPE']
        else:
            row['ID'] = ''
        row['TYPE'] = ''
        if 'SSN' in row.keys():
            row['TEMP SSN'] = row['SSN']
        else:
            row['TEMP SSN'] = ''

    for path in [output_csv_path, system_csv_path]:
      
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        worksheet.append(csv_headers)
        current_subscriber_ssn = None
        subscriber_address = None
        subscriber_city = None
        subscriber_zip = None
        subscriber_state = None
        previous_custodial_parent = None
        for row in csv_data:
            row["PLAN"] = row["PLAN"].zfill(2)
            row["CLASS"] = row["CLASS"].zfill(2)
            for key in row.keys():
                row[key] = str(row[key]) if row[key] is not None else ""

            if 'SUB/DEP' in row.keys():
                if row['SUB/DEP'] != 'Subscriber':
                    row['DEP FIRST NAME'] = str(row.get('FIRST NAME', "")).ljust(20)
                    row['DEP LAST NAME'] = str(row.get('LAST NAME', "")).ljust(20)
                    row['DEP DOB'] = str(row.get('DOB', "")).ljust(20)
                    row['DEP SSN'] = str(row.get('TEMP SSN', "")).ljust(20)
                    row['DEP SEX'] = str(row.get('SEX', "")).ljust(20)

            if 'SEX' in row.keys():
                if row['SEX'] == 'M' and row['SUB/DEP'] == 'Child':
                    row['SUB/DEP'] = 'SON'.ljust(20)
                elif row['SEX'] == 'F' and row['SUB/DEP'] == 'Child':
                    row['SUB/DEP'] = 'DAUGHTER'.ljust(20)
            if 'SUB/DEP' in row.keys():
                if row['SUB/DEP'] == 'Dependent':
                    row['SUB/DEP'] = 'OTHER'.ljust(20)

                if row["SUB/DEP"] == "Subscriber":
                    current_subscriber_ssn = row["SSN"]
                else:
                    row["SSN"] = current_subscriber_ssn
                if row["SUB/DEP"] == "Subscriber":
                    if "ADDRESS 1" in row and row["ADDRESS 1"]:
                        subscriber_address = row["ADDRESS 1"]
                    if 'ZIP' in row.keys() and 'CITY' in row.keys() and 'STATE' in row.keys():
                        subscriber_zip = row["ZIP"]
                        subscriber_city = row["CITY"]
                        subscriber_state = row["STATE"]
                else:
                    if "ADDRESS 1" in row and row["ADDRESS 1"]:
                        if row["ADDRESS 1"] != subscriber_address:
                            row["DEP ADDRESS"] = row["ADDRESS 1"]
                            row["ADDRESS 1"] = subscriber_address
                    if 'ZIP' in row.keys():    
                        if row["ZIP"] != subscriber_zip:
                                row["DEP ZIP"] = row["ZIP"]
                                row["ZIP"] = subscriber_zip
                    if 'CITY' in row.keys():
                        if row["CITY"] != subscriber_city:
                                row["DEP CITY"] = row["CITY"]
                                row["CITY"] = subscriber_city
                    if 'STATE' in row.keys():    
                        if row["STATE"] != subscriber_state:
                                row["DEP STATE"] = row["STATE"]
                                row["STATE"] = subscriber_state                            

            worksheet.append([row.get(header, "") for header in csv_headers])
        

    workbook.save(path)

    cus_df = parse_custodial_data(csv_data)
    total_subscribers = cus_df['SUB/DEP'].str.lower().value_counts().get('subscriber', 0)
    total_spouse_count = cus_df['SUB/DEP'].str.lower().value_counts().get('spouse',0)
    total_dependents = len(cus_df) - total_subscribers
    other_dependents = len(cus_df)-(total_subscribers+total_spouse_count)
    term_check_df = cus_df
    current_date = datetime.now()
    day = current_date.strftime("%A")
    date = current_date.strftime("%Y-%m-%d")
    connection = sqlite3.connect('db.sqlite3')  
    cursor = connection.cursor()
    insert_query = """
    INSERT INTO myapp_count_model (filename, subscriber_count, other_dependent_count, date, day)
    VALUES (?, ?, ?, ?, ?);
    """

    cursor.execute(insert_query, (input_filename, total_subscribers, total_dependents, date, day))

    connection.commit()
    connection.close()

    column_positions = {
    'ADDRESS2': 19,  # Column T
    'TERM DATE': 16,  # Column Q
    'LOCAL': 11,  # Column L
    'EMAIL': 24,  # Column Y
}
    
    def inventory_check(df):
        conn = sqlite3.connect('db.sqlite3')
        inventory_data_query = "SELECT temp_ssn, flag FROM myapp_inventory_table_data"
        inventory_data = pd.read_sql_query(inventory_data_query, conn)
        current_date = datetime.now().strftime('%Y-%m-%d')
        missing_ssn = inventory_data[~inventory_data['temp_ssn'].isin(df['TEMP SSN'])]
        missing_in_inventory = df[~df['TEMP SSN'].isin(inventory_data['temp_ssn'])]
        missing_records = missing_ssn[missing_ssn['flag'] != 'Y']
        missing_ssns = missing_records['temp_ssn'].tolist()
        missing_data = pd.DataFrame()
        try:
            send_missing_ssn_email(missing_ssns,missing_in_inventory['TEMP SSN'].tolist(),'akshay.kumar@onesmarter.com')
            send_missing_ssn_email(missing_ssns,missing_in_inventory['TEMP SSN'].tolist(),'Vikram@vikramsethi.com')
        except:
            print("something went wrong with ssn email")
        print("missing ssn in current file",missing_ssns)
        print("misssing ssn in inventory data",missing_in_inventory['TEMP SSN'])
        if missing_ssns:
            placeholders = ', '.join(['?'] * len(missing_ssns))
            query = f"SELECT * FROM myapp_inventory_table_data WHERE temp_ssn IN ({placeholders})"
            missing_data = pd.read_sql_query(query, conn, params=missing_ssns)
            rename_columns = {
            'last_name': 'LAST NAME',
            'first_name': 'FIRST NAME',
            'ssn': 'SSN',
            'sub_dep': 'SUB/DEP',
            'status': 'STATUS',
            'type': 'TYPE',
            'phone': 'PHONE',
            'address1': 'ADDRESS 1',
            'city': 'CITY',
            'state': 'STATE',
            'zip': 'ZIP',
            'dob': 'DOB',
            'sex': 'SEX',
            'plan': 'PLAN',
            'class_field': 'CLASS',
            'eff_date': 'EFF DATE',
            'id_field': 'ID',
            'dep_first_name': 'DEP FIRST NAME',
            'dep_last_name': 'DEP LAST NAME',
            'dep_dob': 'DEP DOB',
            'dep_ssn': 'DEP SSN',
            'dep_sex': 'DEP SEX',
            'custodial_parent': 'CUSTODIAL PARENT',
            'custodial_address1': 'CUSTODIAL ADDRESS 1',
            'custodial_address2': 'CUSTODIAL ADDRESS 2',
            'custodial_city': 'CUSTODIAL CITY',
            'custodial_state': 'CUSTODIAL STATE',
            'custodial_zip': 'CUSTODIAL ZIP',
            'custodial_phone': 'CUSTODIAL PHONE',
            'address2': 'ADDRESS2',
            'member_id': 'MEMBER ID',
            'date_edi': 'EDI_DATE',
            'filename': 'filename',
            'temp_ssn': 'TEMP SSN',
            'term_date': 'TERM DATE'
        }
            missing_data.rename(columns=rename_columns, inplace=True)
            match = re.search(r'(\d{4})(\d{2})(\d{2})', file_name)
            if match:
                year, month, day = match.groups()
                formatted_date = f"{month}/{day}/{year}"
                print(formatted_date)
            else:
                formatted_date = date
            missing_data['TERM DATE'] = formatted_date
            combined_df = pd.concat([df, missing_data], ignore_index=True)
            update_query = f"UPDATE myapp_inventory_table_data SET flag = 'Y' WHERE temp_ssn IN ({placeholders})"
            conn.execute(update_query, missing_ssns)
            conn.commit()
            conn.close()
        else:
            combined_df = df

 
        if not missing_in_inventory.empty:
            print("init")
            conn = sqlite3.connect("db.sqlite3")
            match = re.search(r'(\d{4})(\d{2})(\d{2})', file_name)
            if match:
                year, month, day = match.groups()
                formatted_date = f"{year}-{month}-{day}"
                print(formatted_date)
            else:
                formatted_date = ''
            cursor = conn.cursor()
            missing_in_inventory['date_edi'] = formatted_date
            data_to_insert = missing_in_inventory.to_dict(orient='records')
            key_to_column = {
                'LAST NAME': 'last_name',
                'FIRST NAME': 'first_name',
                'SSN': 'ssn',
                'SUB/DEP': 'sub_dep',
                'STATUS': 'status',
                'TYPE': 'type',
                'PHONE': 'phone',
                'ADDRESS 1': 'address1',
                'CITY': 'city',
                'STATE': 'state',
                'ZIP': 'zip',
                'DOB': 'dob',
                'SEX': 'sex',
                'PLAN': 'plan',
                'CLASS': 'class_field',
                'EFF DATE': 'eff_date',
                'ID': 'id_field',
                'DEP FIRST NAME': 'dep_first_name',
                'DEP LAST NAME': 'dep_last_name',
                'DEP DOB': 'dep_dob',
                'DEP SSN': 'dep_ssn',
                'DEP SEX': 'dep_sex',
                'CUSTODIAL PARENT': 'custodial_parent',
                'CUSTODIAL ADDRESS 1': 'custodial_address1',
                'CUSTODIAL ADDRESS 2': 'custodial_address2',
                'CUSTODIAL CITY': 'custodial_city',
                'CUSTODIAL STATE': 'custodial_state',
                'CUSTODIAL ZIP': 'custodial_zip',
                'CUSTODIAL PHONE': 'custodial_phone',
                'ADDRESS2': 'address2',
                'MEMBER ID': 'member_id',
                'date_edi': 'date_edi',
                'TEMP SSN': 'temp_ssn',
            }


            db_columns = list(key_to_column.values())

            values = [
                tuple(record.get(key, '').strip() if isinstance(record.get(key), str) else record.get(key)
                    for key in key_to_column.keys())
                for record in data_to_insert
            ]
            conn = sqlite3.connect("db.sqlite3")
            cursor = conn.cursor()

            placeholders = ", ".join(["?"] * len(db_columns)) 
            insert_query = f"""
            INSERT INTO myapp_inventory_table_data ({', '.join(db_columns)})
            VALUES ({placeholders})
            """
            cursor.executemany(insert_query, values)
            conn.commit()
            conn.close()

            print("Missing data inserted successfully!")

        return combined_df

    cus_df = inventory_check(cus_df)
    try:
        cus_df.drop(columns=['TEMP SSN','filename','flag'],inplace=True)
    except:
        pass
    if 'TEMP SSN' in cus_df.columns:
        cus_df.drop(columns=['TEMP SSN'],inplace=True)
    if 'filename' in cus_df.columns:
        cus_df.drop(columns=['filename'],inplace=True)
    if 'TEMP SSN' in cus_df.columns:
        cus_df.drop(columns=['TEMP SSN'],inplace=True)
    if 'flag' in cus_df.columns:
        cus_df.drop(columns=['flag'],inplace=True)
    if 'id' in cus_df.columns:
        cus_df.drop(columns=['id'],inplace=True)
    print(cus_df.columns)


    for col, position in column_positions.items():
        if col not in cus_df.columns:
            cus_df.insert(min(position, len(cus_df.columns)), col, '')
    cus_df.to_csv(output_csv_path)
    print("www",output_xlsx_path)
    try:
        cus_df.to_excel(output_xlsx_path)
    except:
        excel_path = "S:\\OOE\\EDI_PROJECT-\\EDI-Backend\\media\\csv_files"
        new_path = os.path.join(excel_path,f"{input_filename}.xlsx")
        cus_df.to_excel(new_path)
    cus_df['EDI_DATE'] = date
    cus_df['EFF DATE'] = pd.to_datetime(cus_df['EFF DATE'])
    cus_df['EFF DATE'] = cus_df['EFF DATE'].apply(
    lambda x: '01/01/2025' if x < datetime(2025, 1, 1) else x.strftime('%m/%d/%Y')
)
    cus_df_dict = cus_df.to_dict(orient="records")
    DB_NAME = "db.sqlite3"
    def get_previous_business_day():
        today = datetime.today().date()
        if today.weekday() == 0:  # Monday
            return today - timedelta(days=3)  # Friday
        else:
            return today - timedelta(days=1)
    def fetch_data_by_date(date):
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        query = "SELECT member_id FROM myapp_edi_user_data WHERE date_edi = ?"
        cursor.execute(query, (date,))
        data = [row[0] for row in cursor.fetchall()]
        conn.close()
        return set(data)
    
    previous_date = get_previous_business_day().strftime("%Y-%m-%d")
    previous_data = fetch_data_by_date(previous_date)
    if not previous_data:
        print(f"No data available for the previous day: {previous_date}. Exiting process.")
    today_data = set(cus_df['MEMBER ID'].dropna())

    # Find new and missing member IDs
    new_ids = today_data - previous_data
    missing_ids = previous_data - today_data
    # send_member_id_email(new_ids,missing_ids)
    def write_to_queue_django(data, queue_file):
        try:
            with open(queue_file, "ab") as file:
                pickle.dump(data, file)
            print("django input filename added to the queue.")
        except Exception as e:
            print(f"Error writing to the queue: {e}")
    write_to_queue_django(cus_df_dict, "django_queue_file.pkl")
    cus_df.drop(columns=['EDI_DATE'],inplace=True)
    #calculate_count_variance(total_subscribers,total_dependents,input_filename)
    output_outbound_folder = r"S:/OOE/EDI_PROJECT-/EDI-Backend/media/output_outbound"
    file_name = os.path.splitext(os.path.basename(input_file_path))[0]
    output_file_pth = os.path.join(output_outbound_folder, f"{file_name}.834")
    # def write_to_queue_input(data, queue_file):
    #     try:
    #         with open(queue_file, "ab") as file:
    #             pickle.dump(data, file)
    #         print("input filename added to the queue.")
    #     except Exception as e:
    #         print(f"Error writing to the queue: {e}")
    # write_to_queue_input(file_name, "input_queue_file.pkl")
    is_monday = datetime.today().weekday() == 0
    is_first_date = datetime.today().day == 1
    try:
        if is_first_date and is_monday:
            shutil.copy(input_file_path, output_file_pth)
            print(f"File has been saved to: {output_file_pth}")
        elif is_first_date:
            shutil.copy(input_file_path, output_file_pth)
            print(f"File has been saved to: {output_file_pth}")
        elif is_monday:
            shutil.copy(input_file_path, output_file_pth)
            print(f"File has been saved to: {output_file_pth}")
        else:
            print("Today is not monday or first day of the month")
    except Exception as e:
        print(f"Error while saving file: {e}")
        print(f"CSV generated successfully at: {output_csv_path} and {system_csv_path}")

    mapping_dict = {
    'F1': 'SUB/DEP',
    'F2': 'LAST NAME',
    'F3': 'FIRST NAME',
    'F4': 'SSN',
    'F5': 'SEX',
    'F6': 'DOB',
    'F7': 'DEP LAST NAME',
    'F8': 'DEP FIRST NAME',
    'F9': 'DEP SSN',
    'F10': 'DEP SEX',
    'F11': 'DEP DOB',
    'F12': 'LOCAL',
    'F13': 'PLAN',
    'F14': 'CLASS',
    'F15': 'EFF DATE',
    'F16': 'TERM DATE',
    'F17': 'ID',
    'F18': 'ADDRESS 1',
    'F19': 'ADDRESS2',
    'F20': 'CITY',
    'F21': 'STATE',
    'F22': 'ZIP',
    'F23': 'PHONE',
    'F24': 'EMAIL',
    'F25': 'STATUS',
    'F26': 'TYPE',
    'F27': 'MEMBER ID',
    'F28': 'CUSTODIAL PARENT',
    'F29': 'CUSTODIAL ADDRESS 1',
    'F30': 'CUSTODIAL ADDRESS 2',
    'F31': 'CUSTODIAL CITY',
    'F32': 'CUSTODIAL STATE',
    'F33': 'CUSTODIAL ZIP',
    'F34': 'CUSTODIAL PHONE',
}

# Reverse the mapping dictionary to map values to keys
    reverse_mapping = {value: key for key, value in mapping_dict.items()}
    db2_df = cus_df
    term_check_db2_df = db2_df
    extra_rows = term_check_db2_df.merge(term_check_df, how='left', indicator=True).query('_merge == "left_only"')
    print(extra_rows)
    print(type(extra_rows))
    ssn_term_date_list = extra_rows.apply(
        lambda row: {
            'SSN': row['SSN'] if row['SUB/DEP'] == 'Subscriber' else row['DEP SSN'],
            'TERM DATE': row['TERM DATE']
        },
        axis=1
    ).tolist()
    dir  = "media/csv_files"
    filename = "check_term_date.csv"
    path = os.path.join(dir,filename)
    cus_df.to_csv(path)
    try:
        db2_df.drop(columns=['EDI_DATE'],inplace=True)
    except:
        pass
    db2_df['CITY'] = db2_df['CITY'].apply(lambda x: x[:16] if isinstance(x, str) else x)
    print(db2_df.columns)
    print(len(db2_df.columns))
    print('db2 initiated')
    # Rename the columns using the reversed mapping dictionary
    db2_df.rename(columns=reverse_mapping, inplace=True)
    db2_df.fillna(' ',inplace=True)
    column_max_lengths = {
    'F1': 16,
    'F2': 16,
    'F3': 17,
    'F4': 11,
    'F5': 1,
    'F6': 10,
    'F7': 19,
    'F8': 20,
    'F9': 11,
    'F10': 16,
    'F11': 13,
    'F12': 5,
    'F13': 4,
    'F14': 5,
    'F15': 10,
    'F16': 10,
    'F17': 2,
    'F18': 30,
    'F19': 9,
    'F20': 16,
    'F21': 5,
    'F22': 5,
    'F23': 12,
    'F24': 32,
    'F25': 9,
    'F26': 10,
    'F27': 10,
    'F28': 30,
    'F29': 30,
    'F30': 6,
    'F31': 16,
    'F32': 5,
    'F33': 5,
    'F34': 11
}

    for column, max_length in column_max_lengths.items():
        if column in db2_df.columns:
            db2_df[column] = db2_df[column].apply(lambda x: str(x)[:max_length] if pd.notnull(x) else x)
    db2_total_subscribers = db2_df['F1'].str.lower().value_counts().get('subscriber', 0)
    db2_total_spouse_count = db2_df['F1'].str.lower().value_counts().get('spouse',0)
    db2_other_dependents = len(db2_df)-(total_subscribers+total_spouse_count) 
    try:
        if db2_total_subscribers-total_subscribers!=0:
            count_list = [{"excel":[{"total_subscriber":total_subscribers},{"total_spouse":total_spouse_count},{"other dependents":other_dependents}]},{"db2":[{"total_subscriber":db2_total_subscribers},{"total_spouse":db2_total_spouse_count},{"other dependents":db2_other_dependents}]}]
            print(count_list)
            emails = ['krushnarajjadhav015@gmail.com','akshay.kumar@onesmarter.com','Vikram@vikramsethi.com','dprasad@abchldg.com']
            if len(ssn_term_date_list)!= 0:
                send_success_email_count(emails,count_list,ssn_term_date_list)
            else:
                send_success_email_count(emails,count_list,[])
        elif db2_total_spouse_count - total_spouse_count !=0:
            count_list = [{"excel":[{"total_subscriber":total_subscribers},{"total_spouse":total_spouse_count},{"other dependents":other_dependents}]},{"db2":[{"total_subscriber":db2_total_subscribers},{"total_spouse":db2_total_spouse_count},{"other dependents":db2_other_dependents}]}]
            print(count_list)
            emails = ['krushnarajjadhav015@gmail.com','akshay.kumar@onesmarter.com','Vikram@vikramsethi.com','dprasad@abchldg.com']
            if len(ssn_term_date_list)!= 0:
                    send_success_email_count(emails,count_list,ssn_term_date_list)
            else:
                    send_success_email_count(emails,count_list,[])
        elif db2_other_dependents - db2_other_dependents !=0:
            count_list = [{"excel":[{"total_subscriber":total_subscribers},{"total_spouse":total_spouse_count},{"other dependents":other_dependents}]},{"db2":[{"total_subscriber":db2_total_subscribers},{"total_spouse":db2_total_spouse_count},{"other dependents":db2_other_dependents}]}]
            print(count_list)
            emails = ['krushnarajjadhav015@gmail.com','akshay.kumar@onesmarter.com','Vikram@vikramsethi.com','dprasad@abchldg.com']
            if len(ssn_term_date_list)!= 0:
                send_success_email_count(emails,count_list,ssn_term_date_list)
            else:
                send_success_email_count(emails,count_list,[])
        else:
            count_list = [{"excel":[{"total_subscriber":total_subscribers},{"total_spouse":total_spouse_count},{"other dependents":other_dependents}]},{"db2":[{"total_subscriber":db2_total_subscribers},{"total_spouse":db2_total_spouse_count},{"other dependents":db2_other_dependents}]}]
            print(count_list)
            emails = ['krushnarajjadhav015@gmail.com','akshay.kumar@onesmarter.com','dprasad@abchldg.com']
            if len(ssn_term_date_list)!= 0:
                send_success_email_count(emails,count_list,ssn_term_date_list)
            else:
                send_success_email_count(emails,count_list,[])

    except Exception as e:
        count_list = [{"excel":[{"total_subscriber":total_subscribers},{"total_spouse":total_spouse_count},{"other dependents":other_dependents}]},{"db2":[{"total_subscriber":db2_total_subscribers},{"total_spouse":db2_total_spouse_count},{"other dependents":db2_other_dependents}]}]
        print("got the error",count_list)
        print("Email not send due to",e)
    count_list = [{"excel":[{"total_subscriber":total_subscribers},{"total_spouse":total_spouse_count},{"other dependents":other_dependents}]},{"db2":[{"total_subscriber":db2_total_subscribers},{"total_spouse":db2_total_spouse_count},{"other dependents":db2_other_dependents}]}]
    print('sort','error cauisng',count_list)
    print("wrapping up")
    db2_df['F23'] = db2_df['F23'].replace("None", "")
    ddf_dict = db2_df.to_dict(orient='records')
    def write_filename_to_buffer(filename, buffer_file):
        data = {'filename': filename}
        with open(buffer_file, 'w') as file:
            json.dump(data, file)
        print(f"Filename '{filename}' written to {buffer_file}.")
    write_filename_to_buffer(file_name,'file_buffer.json')
    def write_to_queue(data, queue_file):
        try:
            with open(queue_file, "ab") as file:
                pickle.dump(data, file)
            print("Data added to the db2 queue.")
        except Exception as e:
            print(f"Error writing to the queue: {e}")
    with open("output.txt", "w") as file:
        file.write(json.dumps(ddf_dict, indent=4))
    write_to_queue(ddf_dict, "db2_queue.pkl")
    return output_csv_path,edi_excel_path

def parse_custodial_data(csv_data):
    new_df = pd.DataFrame(csv_data)
    new_df['CUSTODIAL ADDRESS 1'] = ''
    new_df['CUSTODIAL ADDRESS 2'] = ''
    new_df['CUSTODIAL CITY'] = ''
    new_df['CUSTODIAL STATE'] = ''
    new_df['CUSTODIAL ZIP'] = ''
    new_df['CUSTODIAL PHONE'] = ''
    if 'ID' in new_df.columns:
        new_df['ID'] = pd.to_numeric(new_df['ID'], errors='coerce')
        condition = new_df['ID'] == 15
    else:
        new_df['id_field'] = pd.to_numeric(new_df['id_field'], errors='coerce')
        condition = new_df['id_field'] == 15
    new_df.fillna('', inplace=True)

    if 'ADDRESS 1' in new_df.columns:
        new_df.loc[condition, 'CUSTODIAL ADDRESS 1'] = new_df.loc[condition, 'ADDRESS 1']
    elif 'address1' in new_df.columns:
        new_df.loc[condition, 'custodial_address_1'] = new_df.loc[condition,'address1']
    if 'ADDRESS 2' in new_df.columns:
        new_df.loc[condition, 'CUSTODIAL ADDRESS 2'] = new_df.loc[condition, 'ADDRESS 2']
    elif 'address2' in new_df.columns:
        new_df.loc[condition, 'custodial_address_2'] = new_df.loc[condition,'address2']
    if 'CITY' in new_df.columns:
        new_df.loc[condition, 'CUSTODIAL CITY'] = new_df.loc[condition, 'CITY']
    elif 'city' in new_df.columns:
        new_df.loc[condition, 'custodial_city'] = new_df.loc[condition,'city']
    if 'STATE' in new_df.columns:
        new_df.loc[condition, 'CUSTODIAL STATE'] = new_df.loc[condition, 'STATE']
    elif 'state' in new_df.columns:
        new_df.loc[condition, 'custodial_state'] = new_df.loc[condition,'state']
    if 'ZIP' in new_df.columns:
        new_df.loc[condition, 'CUSTODIAL ZIP'] = new_df.loc[condition, 'ZIP']
    elif 'zip' in new_df.columns:
        new_df.loc[condition, 'custodial_zip'] = new_df.loc[condition,'zip']
    if 'PHONE' in new_df.columns:
        new_df.loc[condition, 'CUSTODIAL PHONE'] = new_df.loc[condition, 'PHONE']
    elif 'phone' in new_df.columns:
        new_df.loc[condition, 'custodial_phone'] = new_df.loc[condition,'phone']
    
    sdf = new_df
    sdf_data = sdf.to_dict(orient="records")
    previous_custodial_parent = None
    custodial_parent_column = [row.get("CUSTODIAL PARENT", None) for row in sdf_data]
    shifted_custodial_parent_column = [None] + custodial_parent_column[:-1]

    for row, shifted_value in zip(sdf_data, shifted_custodial_parent_column):
        row["CUSTODIAL PARENT"] = shifted_value

    for row in sdf_data:
        if row["SUB/DEP"] != "Subscriber":
                if not row.get("CUSTODIAL ADDRESS 1") or row.get("DEP ADDRESS"):
                        if row.get("DEP ADDRESS"):
                            row["CUSTODIAL ADDRESS 1"] = row["DEP ADDRESS"]
                        elif row.get("ADDRESS 1"):
                            row["CUSTODIAL ADDRESS 1"] = row["ADDRESS 1"]
                if not row.get("CUSTODIAL ZIP") or row.get("DEP ZIP"):
                        if row.get("DEP ZIP"):
                            row["CUSTODIAL ZIP"] = row["DEP ZIP"]
                        elif row.get("ZIP"):
                            row["CUSTODIAL ZIP"] = row["ZIP"]


                if not row.get("CUSTODIAL CITY") or row.get("DEP CITY"): 
                        if row.get("DEP CITY"):
                            row["CUSTODIAL CITY"] = row["DEP CITY"]
                        elif row.get("CITY"):
                            row["CUSTODIAL CITY"] = row["CITY"]


                if not row.get("CUSTODIAL STATE"):
                        if row.get("DEP STATE"):
                            row["CUSTODIAL STATE"] = row["DEP STATE"]
                        elif row.get("STATE"):
                            row["CUSTODIAL STATE"] = row["STATE"]

    sdf = pd.DataFrame(sdf_data)
    desired_order = [
        "SUB/DEP", "LAST NAME", "FIRST NAME", "SSN", "TEMP SSN", "SEX", "DOB",
        "DEP LAST NAME", "DEP FIRST NAME", "DEP SSN", "DEP SEX", "DEP DOB",
        "CUSTODIAL PARENT", "LOCAL", "PLAN", "CLASS", "EFF DATE", "TERM DATE",
        "ID", "ADDRESS 1", "ADDRESS 2", "CITY", "STATE", "ZIP", "PHONE", 
        "EMAIL", "STATUS", "TYPE", "MEMBER ID"
    ]

    existing_columns = sdf.columns.tolist()
    columns_in_order = [col for col in desired_order if col in existing_columns]
    other_columns = [col for col in existing_columns if col not in desired_order]


    final_column_order = columns_in_order + other_columns

    sdf = sdf[final_column_order]
    sdf.drop(columns=['DEP ADDRESS','DEP ZIP','DEP STATE','DEP CITY'],inplace=True)
    output_folder = "media/csv_files/"
    fil  = "new_custodial_report_EDI_834_12-24-2024.xlsx"
    path = os.path.join(output_folder,fil)
    return sdf

def calculate_count_variance(sub_count,dep_count,input_file):
    current_date = datetime.now()
    previous_business_day = current_date - timedelta(days=1)
    while previous_business_day.weekday() > 4: 
        previous_business_day -= timedelta(days=1)

    current_date_str = current_date.strftime("%Y-%m-%d")
    previous_date_str = previous_business_day.strftime("%Y-%m-%d")

    print("Today is:", current_date.strftime("%A, %Y-%m-%d"))
    print("Previous business day:", previous_business_day.strftime("%A, %Y-%m-%d"))

    connection = sqlite3.connect('db.sqlite3')
    cursor = connection.cursor()

    fetch_query = """
    SELECT subscriber_count, dependent_count
    FROM myapp_count_model
    WHERE date = ?;
    """

    today_counts = True

    cursor.execute(fetch_query, (previous_date_str,))
    previous_counts = cursor.fetchone()

    connection.close()

    if today_counts and previous_counts:
        today_subscribers, today_dependents = sub_count,dep_count
        prev_subscribers, prev_dependents = map(int, previous_counts)

        sub_variance = abs(today_subscribers - prev_subscribers) / prev_subscribers * 100
        dep_variance = abs(today_dependents - prev_dependents) / prev_dependents * 100

        email = 'akshay.kumar@onesmarter.com'
        if sub_variance > 5:  
            send_error_email_variance(email,input_file,"Subscriber Count")
        else:
            pass

        if dep_variance > 5:
           send_error_email_variance(email,input_file,"Dependent Count")
        else:
            pass
    else:
        print("Data for one or both dates is missing!")


# print('done')
# output_folder = "media/csv_files/"
# J=output_folder
# archive_folder = "media/archive/"
# os.makedirs(output_folder, exist_ok=True)
# os.makedirs(archive_folder, exist_ok=True)

# input_file_path = r"C:/Users/91942/Downloads/EDI_834_12-24-2024.X12"
# output_file_path= parse_edi_to_csv(input_file_path, output_folder,J)
